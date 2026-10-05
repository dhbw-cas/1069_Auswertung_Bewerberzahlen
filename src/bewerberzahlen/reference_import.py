"""Read final semester counts without evaluating Excel formulas or importing totals."""

from __future__ import annotations

import math
import re
from dataclasses import dataclass, replace
from io import BytesIO
from typing import Literal

from openpyxl import load_workbook

from .constants import FACHBEREICHE
from .storage import semester_date_range

ReferenceMetric = Literal["BEW", "IMM"]
REFERENCE_SHEETS: dict[ReferenceMetric, str] = {
    "BEW": "Entwicklung Conversion BEW-IMMA",
    "IMM": "Entwicklung Imma je Studienjahr",
}


@dataclass(frozen=True)
class ReferenceCount:
    """One explicit count; fachbereich is assigned in the import preview."""

    semester: str
    studiengang: str
    metric: ReferenceMetric
    count: int
    sheet_name: str
    source_cell: str
    fachbereich: str = ""


def semester_from_reference_header(header: str) -> str:
    """Convert a WiSe/SoSe source header to the application's semester key."""
    label = header.strip()
    winter = re.fullmatch(r"WiSe(\d{2}|\d{4})/(\d{2}|\d{4})", label)
    summer = re.fullmatch(r"SoSe(\d{2}|\d{4})", label)
    if winter:
        year = _full_year(winter.group(1))
        suffix = winter.group(2)
        expected = (year + 1) % 100 if len(suffix) == 2 else year + 1
        if int(suffix) != expected:
            raise ValueError(f"Ungültiges Semester im Spaltenkopf: {header}")
        key = f"WS{year}_{(year + 1) % 100:02d}"
    elif summer:
        key = f"SS{_full_year(summer.group(1))}"
    else:
        raise ValueError(f"Ungültiges Semester im Spaltenkopf: {header}")
    semester_date_range(key)
    return key


def _full_year(value: str) -> int:
    return int(value) + (2000 if len(value) == 2 else 0)


def read_reference_workbook_from_bytes(
    content: bytes,
    filename: str,
    *,
    metrics: tuple[ReferenceMetric, ...] = ("BEW", "IMM"),
) -> list[ReferenceCount]:
    """Read selected source sheets; blanks are absent and explicit zeros are retained.

    Only direct numeric semester cells are supported. Formulas, invalid counts,
    duplicate semester headers and duplicate keys are rejected with their source.
    """
    if not filename.lower().endswith(".xlsx") or not content:
        raise ValueError("Bitte eine nicht leere XLSX-Datei mit Semesterreferenzen hochladen.")
    if (
        not metrics
        or len(set(metrics)) != len(metrics)
        or any(metric not in REFERENCE_SHEETS for metric in metrics)
    ):
        raise ValueError("Bitte BEW und/oder IMM genau einmal auswählen.")
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    values: list[ReferenceCount] = []
    try:
        for metric in metrics:
            title = REFERENCE_SHEETS[metric]
            if title not in workbook.sheetnames:
                raise ValueError(f'Erforderlicher Reiter fehlt: "{title}".')
            sheet = workbook[title]
            header_row = next(sheet.iter_rows(min_row=1, max_row=1), ())
            columns: dict[int, str] = {}
            for index, cell in enumerate(header_row):
                header = str(cell.value or "").strip()
                if metric == "BEW":
                    if not header.startswith("BEW "):
                        continue
                    header = header.removeprefix("BEW ").strip()
                elif not header.startswith(("WiSe", "SoSe")):
                    continue
                semester = semester_from_reference_header(header)
                if semester in columns.values():
                    raise ValueError(f"{title}!{cell.coordinate}: Semester kommt mehrfach vor.")
                columns[index] = semester
            if not columns:
                raise ValueError(f"{title}: Keine Semester-Spalten für {metric} gefunden.")
            for row in sheet.iter_rows(min_row=2):
                program = str(row[0].value or "").strip()
                if program == "Gesamtsumme":
                    break
                if not program or program.startswith(("Bereich ", "Fachbereich ")):
                    continue
                for index, semester in columns.items():
                    cell = row[index]
                    value = cell.value
                    if value is None or (isinstance(value, str) and not value.strip()):
                        continue
                    source = f"{title}!{cell.coordinate}"
                    if isinstance(value, bool) or not isinstance(value, (int, float)):
                        raise ValueError(
                            f"{source}: Erwartet wird eine direkte nichtnegative Zahl."
                        )
                    if not math.isfinite(value) or value < 0 or value != int(value):
                        raise ValueError(f"{source}: Erwartet wird eine nichtnegative Ganzzahl.")
                    if value > 2_147_483_647:
                        raise ValueError(f"{source}: Anzahl ist zu groß.")
                    values.append(
                        ReferenceCount(
                            semester, program, metric, int(value), title, cell.coordinate
                        )
                    )
    finally:
        workbook.close()
    if not values:
        raise ValueError("Keine expliziten Semesterzahlen in den gewählten Reitern gefunden.")
    validate_reference_counts(values, require_assignments=False)
    return values


def assign_reference_programs(
    values: list[ReferenceCount], assignments: dict[str, tuple[str, str]]
) -> list[ReferenceCount]:
    """Apply explicit (target program, department) assignments and reject collisions."""
    assigned: list[ReferenceCount] = []
    for value in values:
        if value.studiengang not in assignments:
            raise ValueError(f"Studiengang noch nicht zugeordnet: {value.studiengang}")
        program, department = assignments[value.studiengang]
        assigned.append(replace(value, studiengang=program.strip(), fachbereich=department.strip()))
    validate_reference_counts(assigned)
    return assigned


def validate_reference_counts(
    values: list[ReferenceCount], *, require_assignments: bool = True
) -> None:
    """Validate counts before database writes, including keys after name assignment."""
    if not values:
        raise ValueError("Keine Semesterreferenzen ausgewählt.")
    keys: set[tuple[str, str, str, str]] = set()
    for value in values:
        semester_date_range(value.semester)
        if value.metric not in REFERENCE_SHEETS:
            raise ValueError(f"Ungültige Referenzart: {value.metric}")
        if type(value.count) is not int or not 0 <= value.count <= 2_147_483_647:
            raise ValueError(f"Ungültige Anzahl: {value.sheet_name}!{value.source_cell}")
        if not value.studiengang.strip() or value.studiengang != value.studiengang.strip():
            raise ValueError("Studiengang muss ausgefüllt und ohne äußere Leerzeichen sein.")
        if not value.sheet_name.strip() or not value.source_cell.strip():
            raise ValueError("Quellreiter und Quellzelle müssen ausgefüllt sein.")
        if require_assignments and value.fachbereich not in FACHBEREICHE:
            raise ValueError(f"Fachbereich fehlt oder ist ungültig: {value.studiengang}")
        key = (value.metric, value.semester, value.fachbereich, value.studiengang)
        if key in keys:
            raise ValueError(
                f"Doppelte Referenz: {value.metric}, {value.semester}, {value.studiengang} "
                f"({value.sheet_name}!{value.source_cell})."
            )
        keys.add(key)
